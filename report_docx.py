"""Word (.docx) rendering of a due-diligence report.

Content comes from `report_document.build_report_blocks()`, the same function
`report_pdf.py` and `report_markdown.py` render, so the three exports cannot
disagree about what the analysis found. This file is only the python-docx half.

Word matters here for a specific reason rather than as a checkbox: an IC memo
gets edited before it is circulated, and a PDF cannot be. A VC who wants to add
their own note next to a red flag needs the .docx.
"""
from __future__ import annotations

import io
import re
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from report_document import build_report_blocks

MUTED = RGBColor(0x94, 0xA3, 0xB8)
META = RGBColor(0x4A, 0x55, 0x68)

# The memo keeps its Markdown emphasis (see report_document): Word can render
# it properly as a bold run, where printing the asterisks would just look like
# the export had failed.
_BOLD = re.compile(r"\*\*(.+?)\*\*", re.DOTALL)


def _write_markup(paragraph, text: str) -> None:
    position = 0
    for match in _BOLD.finditer(str(text)):
        if match.start() > position:
            paragraph.add_run(str(text)[position:match.start()])
        paragraph.add_run(match.group(1)).bold = True
        position = match.end()
    remainder = str(text)[position:]
    if remainder or position == 0:
        paragraph.add_run(remainder)


def build_report_docx(report: dict[str, Any], generated_on: str) -> bytes:
    """Word document bytes for a report dict shaped like the API's
    DiligenceResponse (or the equivalent raw_output)."""
    document = Document()

    # Body text a shade smaller than Word's 11pt default, to match the density
    # of the PDF export rather than producing a visibly different document.
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    for block in build_report_blocks(report, generated_on):
        kind = block["kind"]
        if kind == "title":
            document.add_heading(block["text"], level=0)
        elif kind == "meta":
            paragraph = document.add_paragraph()
            run = paragraph.add_run(block["text"])
            run.font.size = Pt(9)
            run.font.color.rgb = META
        elif kind == "heading":
            document.add_heading(block["text"], level=1)
        elif kind == "subheading":
            document.add_heading(re.sub(r"\*\*", "", str(block["text"])), level=2)
        elif kind == "text":
            # Blank lines in an LLM memo are paragraph breaks; emitting the
            # whole memo as one run would collapse its structure.
            for line in str(block["text"]).split("\n"):
                if line.strip():
                    _write_markup(document.add_paragraph(), line)
                else:
                    document.add_paragraph()
        elif kind == "bullets":
            for item in block["items"] or ["None identified."]:
                _write_markup(document.add_paragraph(style="List Bullet"), item)
        elif kind == "table":
            table = document.add_table(rows=1, cols=len(block["header"]))
            table.style = "Light Grid Accent 1"
            for cell, text in zip(table.rows[0].cells, block["header"]):
                cell.text = str(text)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
            for row in block["rows"]:
                cells = table.add_row().cells
                for cell, text in zip(cells, row):
                    cell.text = str(text)
        elif kind == "caveat":
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(block["text"])
            run.font.size = Pt(8)
            run.font.color.rgb = MUTED

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
