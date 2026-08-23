"""Build the sample decks the format tests run against.

The fixtures are committed, so the tests do not depend on running this. It is
kept next to them so the content is inspectable and regenerable rather than
being four opaque binaries nobody can audit.

Why generated rather than hand-collected: the tests need the *same* deck in
four formats to assert that every reader recovers the same facts, and no real
deck exists in all four. They are still realistic rather than stubs -- a
genuine three-column problem slide, a three-column team slide with names,
roles and biographies in separate text boxes, and a pricing table whose
tier/price/limit rows are only correct if reading order is reconstructed
properly. Those are exactly the layouts that were mangling real decks, and a
one-paragraph "hello world" fixture would pass a broken extractor.

    python tests/fixtures/build_fixtures.py
"""
from __future__ import annotations

import sys
from pathlib import Path

FIXTURES = Path(__file__).resolve().parent
sys.path.insert(0, str(FIXTURES.parents[1]))

COMPANY = "Thornbury Freight"
TAGLINE = "Predictive load matching for regional trucking fleets."

PROBLEM_COLUMNS = [
    ("Empty miles", "Regional carriers run 28% of their miles empty because loads are matched by phone."),
    ("Slow quoting", "Brokers take four hours on average to return a rate, so shippers post to three brokers at once."),
    ("Driver churn", "Unpredictable routes drive 91% annual driver turnover at carriers under 100 trucks."),
]

TEAM = [
    ("Priya Raghunathan", "Co-founder & CEO",
     "Ran regional operations for a 400-truck carrier; built its first load-matching desk."),
    ("Marcus Oyelaran", "Co-founder & CTO",
     "Optimisation engineer; previously shipped routing infrastructure at a freight marketplace."),
    ("Helena Vasquez", "Head of Revenue",
     "Sold TMS software to mid-market carriers for six years before joining."),
]

PRICING = [
    ("Owner-Operator", "$49/truck/mo", "Up to 5 trucks"),
    ("Fleet", "$39/truck/mo", "Up to 80 trucks"),
    ("Enterprise", "Custom", "Unlimited trucks"),
]

TRACTION = "We booked $1.4M in gross merchandise value across 62 carriers in our first eight months."
MARKET = "The US regional freight brokerage market is $88B, and we address the $4.2B spent on load matching."


def build_pdf(path: Path) -> None:
    """A multi-column deck, laid out with absolute positioning.

    Drawn with the low-level canvas rather than platypus flowables on purpose:
    platypus lays text out in a single column, which would produce a fixture
    that cannot exercise the column reconstruction it exists to test.
    """
    from reportlab.lib.pagesizes import landscape, letter
    from reportlab.pdfgen import canvas

    width, height = landscape(letter)
    pdf = canvas.Canvas(str(path), pagesize=landscape(letter))

    def wrap(text: str, chars: int) -> list[str]:
        words, lines, current = text.split(), [], ""
        for word in words:
            if len(current) + len(word) + 1 > chars:
                lines.append(current)
                current = word
            else:
                current = f"{current} {word}".strip()
        if current:
            lines.append(current)
        return lines

    def columns(title: str, entries: list[tuple[str, ...]], page_label: str) -> None:
        pdf.setFont("Helvetica-Bold", 26)
        pdf.drawString(50, height - 70, title)
        # Three columns with wide gutters, which is what the reader detects.
        lefts = [60, 300, 540]
        for left, entry in zip(lefts, entries):
            y = height - 150
            pdf.setFont("Helvetica-Bold", 15)
            pdf.drawString(left, y, entry[0])
            y -= 24
            pdf.setFont("Helvetica", 10)
            for field in entry[1:]:
                for line in wrap(field, 34):
                    pdf.drawString(left, y, line)
                    y -= 14
                y -= 6
        pdf.setFont("Helvetica", 9)
        pdf.drawString(50, 40, COMPANY)
        pdf.drawRightString(width - 50, 40, page_label)
        pdf.showPage()

    pdf.setFont("Helvetica-Bold", 34)
    pdf.drawString(60, height / 2 + 20, COMPANY)
    pdf.setFont("Helvetica", 14)
    pdf.drawString(60, height / 2 - 14, TAGLINE)
    pdf.drawString(60, height / 2 - 36, "Logistics - Seed Round Pitch")
    pdf.showPage()

    columns("The Problem", [(h, b) for h, b in PROBLEM_COLUMNS], "2")
    columns("Business Model", [(t, p, limit) for t, p, limit in PRICING], "3")

    pdf.setFont("Helvetica-Bold", 26)
    pdf.drawString(50, height - 70, "Traction & Market")
    pdf.setFont("Helvetica", 12)
    for index, line in enumerate(wrap(TRACTION, 95) + [""] + wrap(MARKET, 95)):
        pdf.drawString(60, height - 130 - index * 18, line)
    pdf.setFont("Helvetica", 9)
    pdf.drawString(50, 40, COMPANY)
    pdf.drawRightString(width - 50, 40, "4")
    pdf.showPage()

    columns("Team", [(n, r, b) for n, r, b in TEAM], "5")
    pdf.save()


def build_pptx(path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Emu, Inches, Pt

    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    blank = presentation.slide_layouts[6]

    def textbox(slide, left, top, width, height, text, size, bold=False):
        box = slide.shapes.add_textbox(Emu(int(left)), Emu(int(top)), Emu(int(width)), Emu(int(height)))
        frame = box.text_frame
        frame.word_wrap = True
        run = frame.paragraphs[0].add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold
        return box

    def column_slide(title, entries, label):
        slide = presentation.slides.add_slide(blank)
        textbox(slide, Inches(0.5), Inches(0.4), Inches(9), Inches(0.8), title, 30, bold=True)
        # Three well-separated columns: each is 3.6in wide with a 0.6in gutter,
        # so the shape-centre clustering in document_extractor has real gaps.
        for index, entry in enumerate(entries):
            left = Inches(0.5) + index * Inches(4.2)
            top = Inches(1.8)
            textbox(slide, left, top, Inches(3.6), Inches(0.5), entry[0], 17, bold=True)
            for offset, field in enumerate(entry[1:], start=1):
                textbox(slide, left, top + Inches(0.65) * offset, Inches(3.6), Inches(1.2), field, 11)
        textbox(slide, Inches(0.5), Inches(6.9), Inches(4), Inches(0.4), f"{COMPANY}  {label}", 9)

    title_slide = presentation.slides.add_slide(blank)
    textbox(title_slide, Inches(0.8), Inches(2.6), Inches(11), Inches(1.0), COMPANY, 40, bold=True)
    textbox(title_slide, Inches(0.8), Inches(3.7), Inches(11), Inches(0.6), TAGLINE, 16)

    column_slide("The Problem", PROBLEM_COLUMNS, "2")
    column_slide("Business Model", PRICING, "3")

    traction = presentation.slides.add_slide(blank)
    textbox(traction, Inches(0.5), Inches(0.4), Inches(9), Inches(0.8), "Traction & Market", 30, bold=True)
    textbox(traction, Inches(0.6), Inches(1.8), Inches(11), Inches(1.0), TRACTION, 13)
    textbox(traction, Inches(0.6), Inches(3.0), Inches(11), Inches(1.0), MARKET, 13)

    column_slide("Team", TEAM, "5")
    presentation.save(str(path))


def build_docx(path: Path) -> None:
    import docx

    document = docx.Document()
    document.add_heading(COMPANY, level=0)
    document.add_paragraph(TAGLINE)

    document.add_heading("The Problem", level=1)
    for heading, body in PROBLEM_COLUMNS:
        document.add_paragraph(heading, style="List Bullet")
        document.add_paragraph(body)

    document.add_heading("Business Model", level=1)
    table = document.add_table(rows=1, cols=3)
    for cell, text in zip(table.rows[0].cells, ("Tier", "Price", "Limit")):
        cell.text = text
    for tier, price, limit in PRICING:
        cells = table.add_row().cells
        cells[0].text, cells[1].text, cells[2].text = tier, price, limit

    document.add_heading("Traction & Market", level=1)
    document.add_paragraph(TRACTION)
    document.add_paragraph(MARKET)

    document.add_heading("Team", level=1)
    for name, role, background in TEAM:
        document.add_paragraph(name)
        document.add_paragraph(role)
        document.add_paragraph(background)

    document.save(str(path))


def build_markdown(path: Path) -> None:
    lines = [f"# {COMPANY}", "", TAGLINE, "", "## The Problem", ""]
    for heading, body in PROBLEM_COLUMNS:
        lines += [heading, body, ""]
    lines += ["## Business Model", "", "| Tier | Price | Limit |", "| --- | --- | --- |"]
    lines += [f"| {t} | {p} | {limit} |" for t, p, limit in PRICING]
    lines += ["", "## Traction & Market", "", TRACTION, "", MARKET, "", "## Team", ""]
    for name, role, background in TEAM:
        lines += [name, role, background, ""]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    build_pdf(FIXTURES / "sample_deck.pdf")
    build_pptx(FIXTURES / "sample_deck.pptx")
    build_docx(FIXTURES / "sample_deck.docx")
    build_markdown(FIXTURES / "sample_deck.md")
    for name in ("sample_deck.pdf", "sample_deck.pptx", "sample_deck.docx", "sample_deck.md"):
        print(f"  wrote {name} ({(FIXTURES / name).stat().st_size:,} bytes)")


if __name__ == "__main__":
    main()
